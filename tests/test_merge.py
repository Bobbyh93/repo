"""WP-1 merge exit criteria, executable.

Runs the gate + adapter + canonical renderer on three inputs and checks:
  1. gate demo (schema v1.2)          -> renders, draft-only, no blockers
  2. migrated legacy chapter 1 (v1.0) -> renders, blocked, _DRAFT stamped
  3. original renderer demo (legacy)  -> renders through the gate, all 26
                                          renderer archetypes survive
Every deck must open with python-pptx with a slide count equal to the
manifest's active slide count.

    pytest -q tests/
"""
from __future__ import annotations

import json
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "skills" / "harrity-lesson-builder-pipeline" / "scripts"
FIXTURES = ROOT / "fixtures"
sys.path.insert(0, str(SCRIPTS))

import generate_lesson_package as renderer  # noqa: E402
import spec_adapter  # noqa: E402
import validate_and_gate as gate  # noqa: E402

from pptx import Presentation  # noqa: E402


def _open_package(outdir: Path):
    manifest = json.loads((outdir / "lesson_manifest.json").read_text(encoding="utf-8"))
    decks = list(outdir.glob("*.pptx"))
    assert len(decks) == 1, decks
    assert zipfile.ZipFile(decks[0]).testzip() is None
    prs = Presentation(str(decks[0]))
    active = [s for s in manifest["slides"] if not s.get("retired")]
    assert len(prs.slides) == len(active)
    for f in manifest["files"]:
        assert (outdir / f).exists(), f
    for name in ("facilitator_guide.md", "learner_handout.md", "assessment_map.csv",
                 "traceability_matrix.csv", "qa_log.md"):
        assert (outdir / name).exists()
    return manifest, decks[0], prs


def _unified_validator(outdir: Path) -> int:
    return subprocess.call([sys.executable, str(SCRIPTS / "validate_unified_package.py"), str(outdir)])


def test_archetype_union_covers_both_sides():
    assert gate.V11_ARCHETYPES <= gate.ARCHETYPES
    assert set(renderer.RENDERERS) <= gate.ARCHETYPES
    # every renderer type is reachable from the adapter
    for t in renderer.RENDERERS:
        assert t in spec_adapter.RENDERER_CONTENT_KEYS, t


def test_case1_gate_demo(tmp_path):
    out = tmp_path / "demo"
    result = gate.run(gate.demo_spec(), out, demo=True)
    assert result["exit"] == 0
    manifest, deck, prs = _open_package(out)
    assert deck.name.startswith("DEMO_") and "_DRAFT" not in deck.name
    assert manifest["qa"]["release_status"] == "draft-only"
    assert manifest["qa"]["defect_counts"]["blocker"] == 0
    assert manifest["governance"]["envelope"] == "master-lesson-1.0.0"
    # five-channel separation: every slide carries its script in the notes slide
    for s in prs.slides:
        assert s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip()
    # v1.1 card_data shapes reach the renderer: MCQ correct answer in notes
    mcq = [s for s in prs.slides][6]
    assert "CORRECT: B" in mcq.notes_slide.notes_text_frame.text
    assert _unified_validator(out) == 0


def test_case2_migrated_ch1_blocks_and_stamps_draft(tmp_path):
    spec = json.loads((FIXTURES / "ch1_migrated_lesson_spec.json").read_text(encoding="utf-8"))
    out = tmp_path / "ch1"
    result = gate.run(spec, out, demo=False)
    assert result["exit"] == 1 and result["blocked"]
    manifest, deck, prs = _open_package(out)
    assert "_DRAFT" in deck.name
    assert manifest["qa"]["release_status"] == "blocked"
    notes = [d["note"] for d in result["defects"] if d["severity"] == "blocker"]
    assert any("organizing_clinical_question" in n for n in notes)
    assert len(prs.slides) == 18
    assert _unified_validator(out) == 0


def test_case3_original_demo_through_gate(tmp_path):
    legacy = json.loads((FIXTURES / "renderer_legacy_demo_spec.json").read_text(encoding="utf-8"))
    spec = spec_adapter.legacy_to_spec(legacy)
    out = tmp_path / "legacy"
    result = gate.run(spec, out, demo=True)
    assert result["exit"] == 0, [d for d in result["defects"] if d["severity"] == "blocker"]
    manifest, deck, prs = _open_package(out)
    assert len(prs.slides) == 26
    kinds = {s["slide_archetype"] for s in manifest["slides"]}
    assert kinds == set(renderer.RENDERERS) - {"generic"}, kinds
    # the gate must flag what the original never checked: no verbatim scripts
    majors = [d["note"] for d in result["defects"] if d["severity"] == "major"]
    assert any("speaker_script under 20 words" in n for n in majors)
    assert _unified_validator(out) == 0


def test_governance_downgrades_unapproved_release_ready(tmp_path):
    spec = gate.demo_spec()
    spec["qa"]["release_status"] = "release-ready"
    result = gate.run(spec, tmp_path / "gov", demo=False)
    assert result["status"] == "faculty-review-needed"
    assert any("release-ready claimed without approvals" in d["note"] for d in result["defects"])


def _fully_approved_demo():
    spec = gate.demo_spec()
    spec["qa"]["release_status"] = "release-ready"
    spec["governance"] = {"promotion_state": "release_ready",
                          "approvals": {k: True for k in gate.APPROVAL_KEYS},
                          "taxonomy_lock": {"status": "locked", "approved_by": "faculty", "approval_date": "2026-09-06"}}
    return spec


def test_governance_release_ready_with_full_approvals(tmp_path):
    spec = _fully_approved_demo()
    spec["qa"]["gates_passed"] = ["visual_qa"]
    result = gate.run(spec, tmp_path / "gov2", demo=False)
    assert result["status"] == "release-ready", result["defects"]


def test_d1_signoffs_alone_do_not_substitute_for_the_visual_gate(tmp_path):
    """D1: seven approvals and the taxonomy lock are human sign-offs; none of
    them is evidence the deck was looked at. Before the fix, release-ready
    stood on approvals alone and `qa.gates_passed` was never consulted, so a
    package no one had rendered could certify as released."""
    spec = _fully_approved_demo()                    # no gates_passed at all
    result = gate.run(spec, tmp_path / "d1a", demo=False)
    assert result["status"] == "faculty-review-needed"
    assert any("required QA gate" in d["note"] for d in result["defects"])


def test_d1_misspelled_gate_name_does_not_satisfy_the_requirement(tmp_path):
    """D1/D3: the requirement is satisfied by the gate that ran, not by any
    string sitting in gates_passed. A typo must leave the gate unmet."""
    spec = _fully_approved_demo()
    spec["qa"]["gates_passed"] = ["visaul_qa", "totally_made_up_gate"]
    result = gate.run(spec, tmp_path / "d1b", demo=False)
    assert result["status"] == "faculty-review-needed"
    assert any("required QA gate" in d["note"] for d in result["defects"])
    assert any("unknown gate 'visaul_qa'" in d["note"] for d in result["defects"])


def test_evidence_gate_blocks_unsourced_claim(tmp_path):
    spec = gate.demo_spec()
    spec["slides"][3]["evidence_status"] = "source-grounded"  # source_refs is empty
    result = gate.run(spec, tmp_path / "ev", demo=False)
    assert result["blocked"]
    assert any("source-grounded with empty source_refs" in d["note"] for d in result["defects"])
    assert "_DRAFT" in Path(result["deck"]).name


def test_renderer_own_demo_still_runs(tmp_path):
    out = tmp_path / "orig"
    files = renderer.generate_package(renderer.make_demo_spec(), out)
    prs = Presentation(files["deck"])
    assert len(prs.slides) == 26


# --- D2: the package validator must not pass on an absent subject -----------
# The suite varied the lesson and always handed the validator a real package.
# It never handed it a manifest that described nothing, which is the state a
# half-written or truncated build actually leaves on disk.

def _validator_output(outdir: Path) -> tuple[int, str]:
    r = subprocess.run([sys.executable, str(SCRIPTS / "validate_unified_package.py"), str(outdir)],
                       capture_output=True, text=True)
    return r.returncode, r.stdout


def test_d2_zero_slide_manifest_cannot_certify_a_deck(tmp_path):
    """D2: `if slides:` made the deck/manifest comparison vacuous. A manifest
    listing no slides passed against a 20-slide deck — the validator's central
    check tested nothing and said so to no one."""
    out = tmp_path / "d2a"
    gate.run(gate.demo_spec(), out, demo=False)
    m = json.loads((out / "lesson_manifest.json").read_text(encoding="utf-8"))
    assert len(m["slides"]) > 0                       # the deck really has slides
    m["slides"] = []
    (out / "lesson_manifest.json").write_text(json.dumps(m), encoding="utf-8")
    code, stdout = _validator_output(out)
    assert code == 1, stdout
    assert "manifest lists no slides" in stdout


def test_d2_empty_files_list_cannot_certify_a_package(tmp_path):
    """D2: an empty files[] made the existence loop vacuous, so a package
    missing every artifact it claimed passed the artifact check."""
    out = tmp_path / "d2b"
    gate.run(gate.demo_spec(), out, demo=False)
    m = json.loads((out / "lesson_manifest.json").read_text(encoding="utf-8"))
    m["files"] = []
    (out / "lesson_manifest.json").write_text(json.dumps(m), encoding="utf-8")
    code, stdout = _validator_output(out)
    assert code == 1, stdout
    assert "claims to contain nothing" in stdout


def test_d4_cjm_rationale_cannot_stand_in_for_the_mapping(tmp_path):
    """D4: any non-empty qa.cjm_coverage_rationale used to excuse every missing
    CJM function at once, including all six. Prose about coverage is not
    coverage; the rationale may only excuse the functions it actually names."""
    spec = gate.demo_spec()
    for s in spec["slides"]:
        s["cjm_functions"] = []
    spec.setdefault("qa", {})["cjm_coverage_rationale"] = "Short package; coverage addressed in the unit exam."
    result = gate.run(spec, tmp_path / "d4a", demo=False)
    assert result["blocked"]
    assert any("no slide maps to any CJM function" in d["note"] for d in result["defects"])


def test_d4_rationale_excuses_only_the_functions_it_names(tmp_path):
    spec = gate.demo_spec()
    named = "evaluate outcomes"
    covered = [f for f in gate.CJM if f != named]
    for i, s in enumerate(spec["slides"]):
        s["cjm_functions"] = [covered[i % len(covered)]]
    qa = spec.setdefault("qa", {})
    qa["cjm_coverage_rationale"] = "Evaluate Outcomes is assessed in the following simulation, not in this lesson."
    ok = gate.run(spec, tmp_path / "d4b", demo=False)
    assert not any("CJM functions never covered" in d["note"] for d in ok["defects"]), ok["defects"]
    # the same rationale must not cover a different gap
    qa["cjm_coverage_rationale"] = "Analyze Cues is assessed in the following simulation."
    bad = gate.run(spec, tmp_path / "d4c", demo=False)
    assert any(named.lower() in d["note"].lower() for d in bad["defects"]
               if "CJM functions never covered" in d["note"])
