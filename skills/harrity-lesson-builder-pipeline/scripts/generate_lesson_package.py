#!/usr/bin/env python3
"""
Generate a Harrity-style nursing lesson package from a JSON lesson spec.

Outputs:
- lesson_deck.pptx
- facilitator_guide.md
- learner_handout.md
- assessment_map.csv
- lesson_manifest.json

Usage:
  python generate_lesson_package.py --spec lesson_spec.json --outdir output_folder
  python generate_lesson_package.py --demo --outdir output_folder
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import re
import sys
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


DARK = "#0F172A"
MUTED = "#475569"
LIGHT = "#F6F8FB"
WHITE = "#FFFFFF"
BORDER = "#D7DFEA"
BLUE = "#2563EB"
TEAL = "#0F766E"
GREEN = "#16A34A"
RED = "#DC2626"
PURPLE = "#7C3AED"
ORANGE = "#EA580C"
YELLOW_BG = "#FFF7ED"
PINK_BG = "#FEF2F2"
BLUE_BG = "#EFF6FF"
TEAL_BG = "#ECFDF5"
GREEN_BG = "#F0FDF4"
PURPLE_BG = "#F5F3FF"

LANE_COLORS = {
    "create": BLUE,
    "support": TEAL,
    "protect": GREEN,
    "signal": RED,
    "teach": BLUE,
    "assess": TEAL,
    "escalate": RED,
    "reassure": GREEN,
    "recognize cues": BLUE,
    "analyze cues": TEAL,
    "prioritize hypotheses": ORANGE,
    "generate solutions": PURPLE,
    "take action": RED,
    "evaluate outcomes": GREEN,
}

CJM_STEPS = [
    "recognize cues",
    "analyze cues",
    "prioritize hypotheses",
    "generate solutions",
    "take action",
    "evaluate outcomes",
]


def rgb(hex_value: str) -> RGBColor:
    value = (hex_value or DARK).strip()
    if value.startswith("#"):
        value = value[1:]
    if len(value) == 3:
        value = "".join(ch * 2 for ch in value)
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def inches(value: float):
    return Inches(value)


def safe_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        return "\n".join(str(x) for x in value)
    return str(value)


def normalize_key(value: Any) -> str:
    return re.sub(r"\s+", " ", safe_text(value).strip().lower())


def first_present(*values: Any, default: Any = "") -> Any:
    for value in values:
        if value is not None and value != "":
            return value
    return default


def get_slide_value(slide_def: Dict[str, Any], key: str, default: Any = None) -> Any:
    if key in slide_def:
        return slide_def[key]
    content = slide_def.get("content") or {}
    return content.get(key, default)


def set_fill(shape, fill_color: Optional[str] = WHITE, line_color: Optional[str] = BORDER, line_width: float = 1.0):
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill_color)
    else:
        shape.fill.background()
    if line_color:
        shape.line.color.rgb = rgb(line_color)
        shape.line.width = Pt(line_width)
    else:
        shape.line.color.rgb = rgb(fill_color or WHITE)
        shape.line.width = Pt(0)


def add_rect(slide, x: float, y: float, w: float, h: float, fill: str = WHITE, line: Optional[str] = BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, inches(x), inches(y), inches(w), inches(h))
    set_fill(shape, fill, line)
    return shape


def add_round_rect(slide, x: float, y: float, w: float, h: float, fill: str = WHITE, line: Optional[str] = BORDER, radius: bool = True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, inches(x), inches(y), inches(w), inches(h))
    set_fill(shape, fill, line)
    return shape


def add_textbox(
    slide,
    text: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    font_size: float = 16,
    bold: bool = False,
    color: str = DARK,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(inches(x), inches(y), inches(w), inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = inches(0.03)
    tf.margin_right = inches(0.03)
    tf.margin_top = inches(0.02)
    tf.margin_bottom = inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = safe_text(text)
    run.font.size = Pt(font_size)
    run.font.bold = bool(bold)
    run.font.name = "Aptos"
    run.font.color.rgb = rgb(color)
    return box


def add_label(slide, text: Any, x: float, y: float, w: float, h: float, fill: str = BLUE_BG, color: str = BLUE, font_size: float = 8):
    add_round_rect(slide, x, y, w, h, fill, line=fill)
    return add_textbox(slide, text, x + 0.05, y + 0.03, w - 0.1, h - 0.04, font_size, True, color, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def add_card(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    title: Any,
    body: Any = "",
    accent: str = BLUE,
    fill: str = WHITE,
    title_size: float = 13,
    body_size: float = 9,
    title_color: str = DARK,
):
    add_round_rect(slide, x, y, w, h, fill, BORDER)
    add_rect(slide, x + 0.06, y + 0.16, 0.05, max(0.15, h - 0.32), accent, accent)
    add_textbox(slide, title, x + 0.18, y + 0.12, w - 0.28, 0.32, title_size, True, title_color)
    if body:
        add_textbox(slide, body, x + 0.18, y + 0.48, w - 0.28, h - 0.58, body_size, False, MUTED)


def add_bottom_callout(slide, text: Any, fill: str = TEAL_BG, color: str = TEAL):
    add_round_rect(slide, 1.3, 6.55, 10.7, 0.45, fill, line=color)
    add_textbox(slide, text, 1.45, 6.66, 10.4, 0.23, 8.5, True, color, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def add_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(LIGHT)


def add_header(slide, slide_def: Dict[str, Any], metadata: Dict[str, Any], idx: int):
    module_label = first_present(get_slide_value(slide_def, "module_label"), metadata.get("chapter"), default="lesson")
    title = first_present(get_slide_value(slide_def, "title"), "lesson slide")
    subtitle = get_slide_value(slide_def, "subtitle", "")
    add_textbox(slide, str(module_label).upper(), 0.35, 0.18, 2.2, 0.18, 6.5, True, BLUE)
    add_textbox(slide, title, 0.35, 0.38, 11.8, 0.42, 19.5, True, DARK)
    if subtitle:
        add_textbox(slide, subtitle, 0.37, 0.82, 11.5, 0.22, 8.5, False, MUTED)
    add_textbox(slide, str(idx), 12.95, 7.12, 0.22, 0.15, 6.5, False, MUTED, PP_ALIGN.RIGHT)


def add_footer(slide, slide_def: Dict[str, Any], idx: int):
    cjm = get_slide_value(slide_def, "cjm_steps", []) or []
    if isinstance(cjm, str):
        cjm = [cjm]
    cjm_text = ", ".join(cjm[:3])
    if cjm_text:
        add_textbox(slide, "CJM: " + cjm_text, 0.35, 7.12, 4.4, 0.14, 5.7, False, MUTED)
    add_textbox(slide, "Harrity Lesson Builder", 5.2, 7.12, 3.0, 0.14, 5.7, False, MUTED, PP_ALIGN.CENTER)
    # WP-1 merge: stable slide id and evidence status travel with the slide
    # so reviewers can see provenance without opening the manifest.
    sid = get_slide_value(slide_def, "slide_id", "")
    ev = get_slide_value(slide_def, "evidence_status", "") or get_slide_value(slide_def, "source_status", "")
    tag = "  ·  ".join(str(x) for x in (sid, ev) if x)
    if tag:
        add_textbox(slide, tag, 8.6, 7.12, 4.2, 0.14, 5.7, False, MUTED, PP_ALIGN.RIGHT)


def add_header_footer(slide, slide_def: Dict[str, Any], metadata: Dict[str, Any], idx: int):
    add_background(slide)
    add_header(slide, slide_def, metadata, idx)
    add_footer(slide, slide_def, idx)


def card_color_for(label: Any, fallback: str = BLUE) -> str:
    key = normalize_key(label)
    return LANE_COLORS.get(key, fallback)


def auto_grid(count: int, left: float = 0.55, top: float = 1.25, right: float = 0.55, bottom: float = 0.85, columns: Optional[int] = None):
    if count <= 0:
        return []
    if not columns:
        if count <= 3:
            columns = count
        elif count == 4:
            columns = 4
        elif count <= 6:
            columns = 3
        else:
            columns = 4
    rows = (count + columns - 1) // columns
    total_w = 13.333 - left - right
    total_h = 7.5 - top - bottom
    gap_x = 0.18
    gap_y = 0.18
    w = (total_w - gap_x * (columns - 1)) / columns
    h = min(1.35, (total_h - gap_y * (rows - 1)) / rows)
    coords = []
    for i in range(count):
        r = i // columns
        c = i % columns
        coords.append((left + c * (w + gap_x), top + r * (h + gap_y), w, h))
    return coords


def add_line(slide, x1: float, y1: float, x2: float, y2: float, color: str = BORDER, width: float = 1.5):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, inches(x1), inches(y1), inches(x2), inches(y2))
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    return line


def render_title(slide, slide_def, spec, idx):
    add_background(slide)
    metadata = spec.get("metadata") or {}
    module = first_present(get_slide_value(slide_def, "module_label"), metadata.get("chapter"), "active learning")
    title = first_present(get_slide_value(slide_def, "title"), metadata.get("topic"), "lesson package")
    subtitle = first_present(get_slide_value(slide_def, "subtitle"), metadata.get("subtitle"), "map-first nursing lesson package")
    add_textbox(slide, str(module).upper(), 0.55, 0.50, 2.4, 0.20, 7.2, True, BLUE)
    add_textbox(slide, title, 0.55, 0.90, 8.2, 0.60, 28, True, DARK)
    add_textbox(slide, subtitle, 0.58, 1.60, 8.3, 0.48, 12, False, MUTED)
    goals = get_slide_value(slide_def, "goals", []) or [
        "see the whole map",
        "identify missed patterns",
        "connect biology to bedside action",
    ]
    add_card(slide, 0.55, 2.55, 5.8, 1.15, "clinical question", spec.get("clinical_question", "what should the nurse notice and do first?"), BLUE, WHITE, 13, 9)
    add_card(slide, 0.55, 3.95, 5.8, 1.45, "today is not a vocabulary tour", "you will use the science only when it explains a patient question or a safety decision.", TEAL, WHITE, 13, 9)
    add_round_rect(slide, 9.25, 0.55, 2.15, 2.15, BLUE_BG, BLUE)
    add_textbox(slide, "Nurse lens", 9.55, 1.00, 1.55, 0.30, 13, True, BLUE, PP_ALIGN.CENTER)
    add_textbox(slide, "TEACH\nASSESS\nESCALATE", 9.75, 1.35, 1.2, 0.75, 10, True, DARK, PP_ALIGN.CENTER)
    y = 5.75
    for i, goal in enumerate(goals[:4]):
        add_card(slide, 0.7 + i * 3.05, y, 2.65, 0.72, str(i + 1), goal, [BLUE, TEAL, GREEN, RED][i % 4], WHITE, 12, 8)
    add_footer(slide, slide_def, idx)


def render_opening_case(slide, slide_def, spec, idx):
    add_header_footer(slide, slide_def, spec.get("metadata") or {}, idx)
    prompt = get_slide_value(slide_def, "patient_prompt", "A patient asks whether a new symptom is normal. What should the nurse ask first?")
    know = get_slide_value(slide_def, "know", []) or []
    need = get_slide_value(slide_def, "need", []) or []
    task = get_slide_value(slide_def, "task", "Separate reassurance, assessment, teaching, and urgent cues.")
    add_card(slide, 0.6, 1.2, 5.9, 1.45, "patient says", prompt, ORANGE, YELLOW_BG, 13, 10)
    add_card(slide, 0.6, 3.0, 2.9, 2.0, "what you know so far", "\n".join("- " + safe_text(x) for x in know), BLUE, WHITE, 12, 8.5)
    add_card(slide, 3.75, 3.0, 2.9, 2.0, "what you do not know yet", "\n".join("- " + safe_text(x) for x in need), PURPLE, WHITE, 12, 8.5)
    add_card(slide, 7.25, 1.2, 4.8, 3.85, "your first task", task, RED, WHITE, 14, 10)
    add_bottom_callout(slide, "Before facts: decide what information is missing and what could be unsafe.", PINK_BG, RED)


def render_clinical_question(slide, slide_def, spec, idx):
    add_header_footer(slide, slide_def, spec.get("metadata") or {}, idx)
    question = first_present(get_slide_value(slide_def, "question"), spec.get("clinical_question"), "what should the nurse notice, connect, protect, and do?")
    add_round_rect(slide, 1.6, 1.05, 10.1, 0.85, DARK, DARK)
    add_textbox(slide, question, 1.85, 1.25, 9.6, 0.42, 16, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    lanes = get_slide_value(slide_def, "lanes", None) or spec.get("concept_lanes") or []
    if not lanes:
        lanes = [
            {"label": "normal", "description": "what should happen"},
            {"label": "cue", "description": "what the nurse notices"},
            {"label": "risk", "description": "what can harm first"},
            {"label": "action", "description": "what the nurse does"},
        ]
    coords = auto_grid(len(lanes), left=0.65, top=2.35, right=0.65, bottom=2.4, columns=min(4, max(1, len(lanes))))
    for i, lane in enumerate(lanes):
        label = lane.get("label", "lane") if isinstance(lane, dict) else safe_text(lane)
        body = lane.get("description", "") if isinstance(lane, dict) else ""
        color = card_color_for(label, [BLUE, TEAL, GREEN, RED, PURPLE, ORANGE][i % 6])
        add_card(slide, *coords[i], label.title(), body, color, WHITE, 14, 9)
    add_bottom_callout(slide, "Every concept must answer: what do I teach, what do I assess, what must I escalate?", TEAL_BG, TEAL)


def render_chapter_map(slide, slide_def, spec, idx):
    add_header_footer(slide, slide_def, spec.get("metadata") or {}, idx)
    milestones = get_slide_value(slide_def, "milestones", []) or []
    if not milestones:
        milestones = [{"label": lane.get("label", f"step {i+1}"), "body": lane.get("description", "")} for i, lane in enumerate(spec.get("concept_lanes", []))]
    n = max(1, len(milestones))
    start_x = 0.85
    end_x = 12.15
    y = 2.35
    if n > 1:
        add_line(slide, start_x + 0.4, y, end_x - 0.4, y, BORDER, 3)
    for i, m in enumerate(milestones):
        x = start_x + (end_x - start_x) * i / max(1, n - 1)
        color = [BLUE, TEAL, GREEN, ORANGE, RED, PURPLE][i % 6]
        add_round_rect(slide, x - 0.19, y - 0.19, 0.38, 0.38, color, color)
        add_textbox(slide, str(i + 1), x - 0.12, y - 0.10, 0.25, 0.18, 8, True, WHITE, PP_ALIGN.CENTER)
        add_textbox(slide, safe_text(m.get("label", "step")).title(), x - 0.75, y + 0.42, 1.5, 0.28, 10, True, DARK, PP_ALIGN.CENTER)
        add_textbox(slide, m.get("body", ""), x - 0.8, y + 0.75, 1.6, 0.45, 6.6, False, MUTED, PP_ALIGN.CENTER)
    add_bottom_callout(slide, "Map first. Drill down only where the map or performance data shows a weak link.", TEAL_BG, TEAL)


def render_warmup_sequence(slide, slide_def, spec, idx):
    add_header_footer(slide, slide_def, spec.get("metadata") or {}, idx)
    prompt = get_slide_value(slide_def, "prompt", "Place these events in the most logical sequence.")
    items = get_slide_value(slide_def, "items", []) or []
    add_card(slide, 0.75, 1.10, 11.8, 0.75, "learner task", prompt, PURPLE, WHITE, 13, 10)
    y = 2.15
    for i, item in enumerate(items[:6]):
        add_round_rect(slide, 1.05, y + i * 0.62, 10.95, 0.42, WHITE, BORDER)
        add_textbox(slide, chr(ord("A") + i), 1.2, y + i * 0.71 - 0.03, 0.3, 0.16, 8, True, PURPLE)
        add_textbox(slide, item, 1.6, y + i * 0.61 + 0.08, 10.0, 0.18, 8.5, False, MUTED)
    add_bottom_callout(slide, "Use sequence misses to identify where the concept chain breaks.", PURPLE_BG, PURPLE)


def render_concept_cards(slide, slide_def, spec, idx):
    add_header_footer(slide, slide_def, spec.get("metadata") or {}, idx)
    cards = get_slide_value(slide_def, "cards", []) or []
    coords = auto_grid(len(cards), columns=3 if len(cards) > 4 else len(cards), top=1.25, bottom=1.35)
    for i, card in enumerate(cards):
        label = card.get("title", card.get("label", f"concept {i + 1}"))
        body = card.get("body", card.get("description", ""))
        add_card(slide, *coords[i], label, body, [BLUE, TEAL, GREEN, PURPLE, ORANGE, RED][i % 6], WHITE)
    callout = get_slide_value(slide_def, "callout", "Structure matters because it predicts function, cue, and nursing response.")
    add_bottom_callout(slide, callout, BLUE_BG, BLUE)


def render_match_activity(slide, slide_def, spec, idx):
    add_header_footer(slide, slide_def, spec.get("metadata") or {}, idx)
    sources = get_slide_value(slide_def, "sources", []) or []
    targets = get_slide_value(slide_def, "targets", []) or []
    add_textbox(slide, "Work with one partner: match each source to the cue or clinical meaning most directly involved.", 0.8, 1.10, 11.7, 0.28, 9.5, True, DARK)
    max_rows = max(len(sources), len(targets), 1)
    for i in range(max_rows):
        y = 1.65 + i * 0.76
        if i < len(sources):
            src = sources[i]
            add_card(slide, 0.85, y, 5.25, 0.58, chr(ord("A") + i), safe_text(src), BLUE, WHITE, 10.5, 7.5)
        if i < len(targets):
            tgt = targets[i]
            add_card(slide, 7.15, y, 4.95, 0.58, str(i + 1), safe_text(tgt), TEAL, WHITE, 10.5, 7.5)
    add_bottom_callout(slide, "Output: write matches and one sentence explaining one match.", TEAL_BG, TEAL)


def render_debrief(slide, slide_def, spec, idx):
    add_header_footer(slide, slide_def, spec.get("metadata") or {}, idx)
    cards = get_slide_value(slide_def, "cards", []) or []
    coords = auto_grid(len(cards), columns=2 if len(cards) <= 4 else 3, top=1.2, bottom=1.25)
    for i, card in enumerate(cards):
        title = card.get("title", f"reason {i + 1}")
        body = card.get("body", "")
        add_card(slide, *coords[i], title, body, [BLUE, TEAL, RED, PURPLE, GREEN, ORANGE][i % 6], WHITE, 12, 8.2)
    add_bottom_callout(slide, get_slide_value(slide_def, "callout", "Clinical reasoning improves when learners can explain why the cue matters."), YELLOW_BG, ORANGE)


def render_control_room(slide, slide_def, spec, idx):
    add_header_footer(slide, slide_def, spec.get("metadata") or {}, idx)
    center = get_slide_value(slide_def, "center", "control system")
    nodes = get_slide_value(slide_def, "nodes", []) or []
    add_round_rect(slide, 5.05, 2.15, 3.2, 1.05, DARK, DARK)
    add_textbox(slide, center, 5.28, 2.42, 2.75, 0.45, 16, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    coords = [(0.8, 1.25), (9.3, 1.25), (0.8, 4.35), (9.3, 4.35)]
    for i, node in enumerate(nodes[:4]):
        title = node.get("title", f"node {i + 1}") if isinstance(node, dict) else safe_text(node)
        body = node.get("body", "") if isinstance(node, dict) else ""
        x, y = coords[i]
        add_card(slide, x, y, 3.2, 1.15, title, body, [BLUE, TEAL, PURPLE, GREEN][i % 4], WHITE, 12, 8)
        add_line(slide, x + (3.2 if x < 5 else 0), y + 0.58, 5.05 if x < 5 else 8.25, 2.68, BORDER, 1.3)
    add_bottom_callout(slide, get_slide_value(slide_def, "callout", "Control systems explain timing, symptoms, and safe teaching."), BLUE_BG, BLUE)


def render_timeline(slide, slide_def, spec, idx):
    add_header_footer(slide, slide_def, spec.get("metadata") or {}, idx)
    milestones = get_slide_value(slide_def, "milestones", []) or []
    n = max(1, len(milestones))
    start_x = 1.0
    end_x = 12.0
    y = 2.55
    add_line(slide, start_x, y, end_x, y, BORDER, 3)
    for i, m in enumerate(milestones):
        x = start_x + (end_x - start_x) * i / max(1, n - 1)
        color = [BLUE, TEAL, PURPLE, GREEN, ORANGE, RED][i % 6]
        add_round_rect(slide, x - 0.16, y - 0.16, 0.32, 0.32, color, color)
        add_textbox(slide, m.get("label", f"step {i+1}"), x - 0.8, y + 0.38, 1.6, 0.25, 9.5, True, DARK, PP_ALIGN.CENTER)
        add_textbox(slide, m.get("body", ""), x - 0.85, y + 0.70, 1.7, 0.55, 7, False, MUTED, PP_ALIGN.CENTER)
    add_bottom_callout(slide, get_slide_value(slide_def, "callout", "Timing explains which patient teaching and safety cues matter."), TEAL_BG, TEAL)


def render_checkpoint_mcq(slide, slide_def, spec, idx):
    add_header_footer(slide, slide_def, spec.get("metadata") or {}, idx)
    question = get_slide_value(slide_def, "question", "Which response is safest?")
    options = get_slide_value(slide_def, "options", []) or []
    add_card(slide, 0.85, 1.20, 11.7, 0.95, "checkpoint question", question, PURPLE, PURPLE_BG, 13, 10)
    y = 2.45
    for i, option in enumerate(options[:5]):
        add_round_rect(slide, 1.1, y + i * 0.62, 10.95, 0.44, WHITE, BORDER)
        add_textbox(slide, chr(ord("A") + i), 1.25, y + i * 0.62 + 0.13, 0.28, 0.16, 8, True, PURPLE)
        add_textbox(slide, option, 1.7, y + i * 0.62 + 0.11, 9.9, 0.18, 8.5, False, MUTED)
    add_bottom_callout(slide, get_slide_value(slide_def, "callout", "Ask: what information is missing, and what could be unsafe to promise?"), PURPLE_BG, PURPLE)


def render_process_map(slide, slide_def, spec, idx):
    render_chapter_map(slide, slide_def, spec, idx)


def render_teaching_point(slide, slide_def, spec, idx):
    add_header_footer(slide, slide_def, spec.get("metadata") or {}, idx)
    cards = get_slide_value(slide_def, "cards", []) or []
    coords = auto_grid(len(cards), columns=max(1, min(3, len(cards))), top=1.35, bottom=2.45)
    for i, card in enumerate(cards):
        add_card(slide, *coords[i], card.get("title", f"point {i+1}"), card.get("body", ""), [BLUE, TEAL, GREEN][i % 3], WHITE, 13, 8.5)
    warning = get_slide_value(slide_def, "warning", "Do not overpromise safety when assessment or escalation is needed.")
    add_card(slide, 0.9, 5.25, 11.5, 0.88, "safety teaching limit", warning, ORANGE, YELLOW_BG, 12, 8.5)
    add_bottom_callout(slide, get_slide_value(slide_def, "callout", "Patient teaching is an action: clear, safe, and honest."), TEAL_BG, TEAL)


def render_exchange_model(slide, slide_def, spec, idx):
    add_header_footer(slide, slide_def, spec.get("metadata") or {}, idx)
    left = get_slide_value(slide_def, "left", {"title": "source", "body": "input"})
    center = get_slide_value(slide_def, "center", {"title": "exchange", "body": "organ or process"})
    right = get_slide_value(slide_def, "right", {"title": "recipient", "body": "output"})
    add_round_rect(slide, 1.0, 2.0, 2.25, 2.25, BLUE_BG, BLUE)
    add_textbox(slide, left.get("title", "source"), 1.25, 2.55, 1.75, 0.28, 13, True, BLUE, PP_ALIGN.CENTER)
    add_textbox(slide, left.get("body", ""), 1.25, 2.95, 1.75, 0.5, 8, False, MUTED, PP_ALIGN.CENTER)
    add_round_rect(slide, 5.0, 1.85, 3.35, 2.55, WHITE, BORDER)
    add_textbox(slide, center.get("title", "exchange"), 5.3, 2.35, 2.75, 0.32, 15, True, DARK, PP_ALIGN.CENTER)
    add_textbox(slide, center.get("body", ""), 5.35, 2.80, 2.65, 0.72, 8.5, False, MUTED, PP_ALIGN.CENTER)
    add_round_rect(slide, 10.1, 2.0, 2.25, 2.25, TEAL_BG, TEAL)
    add_textbox(slide, right.get("title", "recipient"), 10.35, 2.55, 1.75, 0.28, 13, True, TEAL, PP_ALIGN.CENTER)
    add_textbox(slide, right.get("body", ""), 10.35, 2.95, 1.75, 0.5, 8, False, MUTED, PP_ALIGN.CENTER)
    add_line(slide, 3.25, 3.12, 5.0, 3.12, BORDER, 2)
    add_line(slide, 8.35, 3.12, 10.1, 3.12, BORDER, 2)
    warning = get_slide_value(slide_def, "warning", "Exchange systems support function, but they are not perfect shields.")
    add_card(slide, 1.35, 5.05, 10.65, 0.75, "nursing misconception to correct", warning, RED, PINK_BG, 12, 8)
    add_bottom_callout(slide, get_slide_value(slide_def, "callout", "Teach the function and the safety limit together."), PINK_BG, RED)


def render_triad(slide, slide_def, spec, idx):
    add_header_footer(slide, slide_def, spec.get("metadata") or {}, idx)
    cards = get_slide_value(slide_def, "cards", []) or []
    coords = auto_grid(len(cards), columns=max(1, min(3, len(cards))), top=1.55, bottom=2.0)
    for i, card in enumerate(cards):
        add_card(slide, *coords[i], card.get("title", f"support {i+1}"), card.get("body", ""), [BLUE, TEAL, PURPLE][i % 3], WHITE, 13, 8.5)
    add_bottom_callout(slide, get_slide_value(slide_def, "callout", "Risk belongs to the patient outcome: not one structure alone."), YELLOW_BG, ORANGE)


def render_safety_chain(slide, slide_def, spec, idx):
    render_chapter_map(slide, slide_def, spec, idx)


def render_compare(slide, slide_def, spec, idx):
    add_header_footer(slide, slide_def, spec.get("metadata") or {}, idx)
    columns = get_slide_value(slide_def, "columns", []) or []
    x_positions = [0.9, 6.85]
    for i, col in enumerate(columns[:2]):
        color = [RED, TEAL][i % 2]
        add_card(slide, x_positions[i], 1.35, 5.55, 3.5, col.get("title", f"column {i+1}"), "\n".join("- " + safe_text(x) for x in col.get("items", [])), color, WHITE, 15, 9)
    add_bottom_callout(slide, get_slide_value(slide_def, "callout", "One-minute prompt: explain this distinction to a patient without using the word 'dangerous' loosely."), PINK_BG, RED)


def render_risk_engine(slide, slide_def, spec, idx):
    add_header_footer(slide, slide_def, spec.get("metadata") or {}, idx)
    center = get_slide_value(slide_def, "center", "risk question")
    factors = get_slide_value(slide_def, "factors", []) or []
    add_round_rect(slide, 5.1, 2.35, 3.1, 1.05, DARK, DARK)
    add_textbox(slide, center, 5.35, 2.62, 2.6, 0.40, 14, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    coords = [(1.1, 1.35), (5.1, 1.05), (9.1, 1.35), (1.1, 4.35), (5.1, 4.75), (9.1, 4.35)]
    for i, factor in enumerate(factors[:6]):
        title = factor.get("title", f"factor {i+1}") if isinstance(factor, dict) else safe_text(factor)
        body = factor.get("body", "") if isinstance(factor, dict) else ""
        x, y = coords[i]
        add_card(slide, x, y, 3.1, 0.9, title, body, [ORANGE, BLUE, TEAL, PURPLE, GREEN, RED][i % 6], WHITE, 11.5, 7.4)
    add_bottom_callout(slide, get_slide_value(slide_def, "callout", "Risk is a pattern: exposure, timing, dose, susceptibility, and clinical context."), PURPLE_BG, PURPLE)


def render_mini_case(slide, slide_def, spec, idx):
    add_header_footer(slide, slide_def, spec.get("metadata") or {}, idx)
    scenario = get_slide_value(slide_def, "scenario", "A patient asks a safety question. Respond without guessing.")
    steps = get_slide_value(slide_def, "steps", []) or []
    add_card(slide, 0.8, 1.15, 11.75, 0.78, "mini case", scenario, ORANGE, YELLOW_BG, 12, 9)
    coords = auto_grid(len(steps), columns=max(1, min(3, len(steps))), top=2.25, bottom=2.25)
    for i, step in enumerate(steps):
        add_card(slide, *coords[i], step.get("title", f"step {i+1}"), step.get("body", ""), [RED, BLUE, PURPLE, TEAL][i % 4], WHITE, 12, 8)
    add_bottom_callout(slide, get_slide_value(slide_def, "callout", "Safe response = assess first, avoid false certainty, give specific next step."), TEAL_BG, TEAL)


def render_urgency_sort(slide, slide_def, spec, idx):
    add_header_footer(slide, slide_def, spec.get("metadata") or {}, idx)
    categories = get_slide_value(slide_def, "categories", []) or []
    coords = auto_grid(len(categories), columns=2 if len(categories) <= 4 else 3, top=1.25, bottom=1.35)
    for i, cat in enumerate(categories):
        title = cat.get("title", f"category {i+1}")
        items = cat.get("items", [])
        add_card(slide, *coords[i], title, "\n".join("- " + safe_text(x) for x in items), [RED, BLUE, TEAL, PURPLE, ORANGE, GREEN][i % 6], WHITE, 12, 8)
    add_bottom_callout(slide, get_slide_value(slide_def, "callout", "Escalation is a nursing action when cues suggest immediate risk."), PINK_BG, RED)


def render_script_template(slide, slide_def, spec, idx):
    add_header_footer(slide, slide_def, spec.get("metadata") or {}, idx)
    steps = get_slide_value(slide_def, "steps", []) or []
    coords = auto_grid(len(steps), columns=3 if len(steps) > 4 else 2, top=1.25, bottom=1.25)
    for i, step in enumerate(steps):
        title = step.get("title", f"step {i+1}")
        body = step.get("body", "")
        add_card(slide, *coords[i], f"{i + 1}. {title}", body, [TEAL, BLUE, PURPLE, GREEN, ORANGE, RED][i % 6], WHITE, 12, 8)
    add_bottom_callout(slide, get_slide_value(slide_def, "callout", "Accurate + plain + nonjudgmental beats technical + vague."), TEAL_BG, TEAL)


def render_basics_grid(slide, slide_def, spec, idx):
    render_concept_cards(slide, slide_def, spec, idx)


def render_capstone_mcq(slide, slide_def, spec, idx):
    add_header_footer(slide, slide_def, spec.get("metadata") or {}, idx)
    scenario = get_slide_value(slide_def, "scenario", "A patient has a new concerning cue.")
    question = get_slide_value(slide_def, "question", "What is the priority nursing response?")
    options = get_slide_value(slide_def, "options", []) or []
    add_card(slide, 0.75, 1.10, 11.85, 0.68, "capstone scenario", scenario, RED, PINK_BG, 11.5, 8.5)
    add_card(slide, 1.0, 2.05, 11.35, 0.72, "priority nursing response", question, PURPLE, WHITE, 12.5, 9)
    y = 3.05
    for i, option in enumerate(options[:5]):
        add_round_rect(slide, 1.2, y + i * 0.58, 10.8, 0.40, WHITE, BORDER)
        add_textbox(slide, chr(ord("A") + i), 1.35, y + i * 0.58 + 0.12, 0.28, 0.16, 8, True, RED)
        add_textbox(slide, option, 1.78, y + i * 0.58 + 0.09, 9.7, 0.18, 8.4, False, MUTED)
    add_bottom_callout(slide, get_slide_value(slide_def, "callout", "Choose the response that stabilizes risk first and keeps assessment moving."), PINK_BG, RED)


def render_debrief_three(slide, slide_def, spec, idx):
    add_header_footer(slide, slide_def, spec.get("metadata") or {}, idx)
    cards = get_slide_value(slide_def, "cards", []) or []
    coords = auto_grid(len(cards), columns=max(1, min(3, len(cards))), top=1.35, bottom=2.2)
    for i, card in enumerate(cards):
        add_card(slide, *coords[i], card.get("title", f"reason {i+1}"), card.get("body", ""), [RED, ORANGE, BLUE, GREEN][i % 4], [PINK_BG, YELLOW_BG, BLUE_BG, GREEN_BG][i % 4], 12, 8)
    add_bottom_callout(slide, get_slide_value(slide_def, "callout", "Teach when stable. Escalate when red flags appear. Reassess understanding and response."), TEAL_BG, TEAL)


def render_retrieval_check(slide, slide_def, spec, idx):
    add_header_footer(slide, slide_def, spec.get("metadata") or {}, idx)
    questions = get_slide_value(slide_def, "questions", []) or []
    y = 1.25
    for i, q in enumerate(questions[:6]):
        add_round_rect(slide, 0.85, y + i * 0.72, 11.65, 0.48, WHITE, BORDER)
        color = [BLUE, PURPLE, TEAL, RED, GREEN, ORANGE][i % 6]
        add_rect(slide, 0.95, y + i * 0.72 + 0.08, 0.05, 0.32, color, color)
        add_textbox(slide, safe_text(q), 1.15, y + i * 0.72 + 0.13, 10.9, 0.18, 8.7, False, DARK)
    add_bottom_callout(slide, get_slide_value(slide_def, "callout", "Answer key goes in the facilitator guide."), WHITE, MUTED)


def render_takeaway(slide, slide_def, spec, idx):
    add_header_footer(slide, slide_def, spec.get("metadata") or {}, idx)
    actions = get_slide_value(slide_def, "actions", []) or [
        {"title": "teach", "body": "use plain language"},
        {"title": "assess", "body": "ask targeted questions"},
        {"title": "escalate", "body": "identify red flags"},
        {"title": "reassure", "body": "only when safe"},
    ]
    coords = auto_grid(len(actions), columns=max(1, min(4, len(actions))), top=1.45, bottom=2.2)
    for i, action in enumerate(actions):
        title = action.get("title", f"action {i+1}")
        body = action.get("body", "")
        add_card(slide, *coords[i], title.title(), body, card_color_for(title, [BLUE, TEAL, RED, GREEN][i % 4]), WHITE, 13, 8.5)
    statement = get_slide_value(slide_def, "statement", "A safe nurse can explain the process, reduce preventable risk, screen for urgent cues, and connect the patient to timely follow-up.")
    add_card(slide, 1.35, 5.35, 10.55, 0.9, "final bedside frame", statement, DARK, WHITE, 12, 8)


def render_generic(slide, slide_def, spec, idx):
    add_header_footer(slide, slide_def, spec.get("metadata") or {}, idx)
    body = get_slide_value(slide_def, "body", "")
    bullets = get_slide_value(slide_def, "bullets", []) or []
    if bullets:
        body = "\n".join("- " + safe_text(x) for x in bullets)
    add_card(slide, 0.95, 1.35, 11.35, 4.65, get_slide_value(slide_def, "lead", "key idea"), body, BLUE, WHITE, 15, 11)
    callout = get_slide_value(slide_def, "callout", "Connect this idea to what the nurse recognizes, analyzes, prioritizes, does, and evaluates.")
    add_bottom_callout(slide, callout, TEAL_BG, TEAL)


RENDERERS = {
    "title": render_title,
    "opening_case": render_opening_case,
    "clinical_question": render_clinical_question,
    "chapter_map": render_chapter_map,
    "warmup_sequence": render_warmup_sequence,
    "concept_cards": render_concept_cards,
    "match_activity": render_match_activity,
    "debrief": render_debrief,
    "control_room": render_control_room,
    "timeline": render_timeline,
    "checkpoint_mcq": render_checkpoint_mcq,
    "process_map": render_process_map,
    "teaching_point": render_teaching_point,
    "exchange_model": render_exchange_model,
    "triad": render_triad,
    "safety_chain": render_safety_chain,
    "compare": render_compare,
    "risk_engine": render_risk_engine,
    "mini_case": render_mini_case,
    "urgency_sort": render_urgency_sort,
    "script_template": render_script_template,
    "basics_grid": render_basics_grid,
    "capstone_mcq": render_capstone_mcq,
    "debrief_three": render_debrief_three,
    "retrieval_check": render_retrieval_check,
    "takeaway": render_takeaway,
    "generic": render_generic,
}


def make_presentation(spec: Dict[str, Any], out_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = inches(13.333)
    prs.slide_height = inches(7.5)
    blank = prs.slide_layouts[6]
    slides = spec.get("slides")
    if not isinstance(slides, list) or not slides:
        raise ValueError("lesson spec must include a non-empty slides array unless --demo is used")
    spec["slides"] = slides
    for idx, slide_def in enumerate(slides, start=1):
        slide = prs.slides.add_slide(blank)
        slide_type = normalize_key(slide_def.get("type", "generic")).replace(" ", "_")
        renderer = RENDERERS.get(slide_type, render_generic)
        renderer(slide, slide_def, spec, idx)
        # WP-1 merge: the narration channel goes to the notes slide, never on
        # the slide body (five-channel separation).
        notes = safe_text(get_slide_value(slide_def, "speaker_notes", ""))
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    prs.save(out_path)


def cjm_coverage(slides: List[Dict[str, Any]]) -> Dict[str, int]:
    coverage = {step: 0 for step in CJM_STEPS}
    for slide in slides:
        steps = get_slide_value(slide, "cjm_steps", []) or []
        if isinstance(steps, str):
            steps = [steps]
        for step in steps:
            key = normalize_key(step)
            if key in coverage:
                coverage[key] += 1
    return coverage


def write_assessment_map(spec: Dict[str, Any], out_path: Path) -> None:
    slides = spec.get("slides") or []
    with out_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow([
            "slide_number",
            "slide_id",
            "slide_title",
            "slide_type",
            "concepts",
            "cjm_steps",
            "learner_task",
            "remediation_target",
            "evidence_of_learning",
        ])
        for i, slide in enumerate(slides, start=1):
            writer.writerow([
                i,
                get_slide_value(slide, "slide_id", f"S{i:02d}"),
                get_slide_value(slide, "title", ""),
                slide.get("type", "generic"),
                "; ".join(safe_text(x) for x in (get_slide_value(slide, "concepts", []) or [])),
                "; ".join(safe_text(x) for x in (get_slide_value(slide, "cjm_steps", []) or [])),
                get_slide_value(slide, "learner_task", ""),
                get_slide_value(slide, "remediation_target", ""),
                get_slide_value(slide, "evidence_of_learning", ""),
            ])


def write_facilitator_guide(spec: Dict[str, Any], out_path: Path) -> None:
    metadata = spec.get("metadata") or {}
    slides = spec.get("slides") or []
    coverage = cjm_coverage(slides)
    lines: List[str] = []
    lines.append(f"# Facilitator Guide: {metadata.get('chapter', 'lesson')} - {metadata.get('topic', 'topic')}\n")
    lines.append(f"**Clinical question:** {spec.get('clinical_question', 'not specified')}\n")
    lines.append("## CJM Coverage\n")
    for step, count in coverage.items():
        lines.append(f"- {step}: {count}")
    lines.append("\n## Teaching Flow\n")
    for i, slide in enumerate(slides, start=1):
        title = get_slide_value(slide, "title", f"slide {i}")
        cjm = ", ".join(safe_text(x) for x in (get_slide_value(slide, "cjm_steps", []) or []))
        learner_task = get_slide_value(slide, "learner_task", "")
        notes = get_slide_value(slide, "speaker_notes", "")
        answer = get_slide_value(slide, "answer", "")
        rationale = get_slide_value(slide, "rationale", "")
        lines.append(f"### Slide {i}: {title}")
        if cjm:
            lines.append(f"- CJM: {cjm}")
        if learner_task:
            lines.append(f"- Learner task: {learner_task}")
        if answer:
            lines.append(f"- Answer key: {answer}")
        if rationale:
            lines.append(f"- Rationale: {rationale}")
        if notes:
            lines.append(f"- Facilitator notes: {notes}")
        lines.append("")
    quiz = spec.get("quiz_analysis") or []
    if quiz:
        lines.append("## Quiz-Based Drill-Down Plan\n")
        for item in quiz:
            lines.append(f"- **{item.get('objective', 'objective')}:** {item.get('score_percent', 'n/a')}% | {item.get('miss_type', 'miss type')} | {item.get('remediation', 'remediation not specified')}")
        lines.append("")
    out_path.write_text("\n".join(lines), encoding="utf-8")


def write_learner_handout(spec: Dict[str, Any], out_path: Path) -> None:
    metadata = spec.get("metadata") or {}
    slides = spec.get("slides") or []
    lines: List[str] = []
    lines.append(f"# Learner Handout: {metadata.get('topic', 'lesson')}\n")
    lines.append(f"**Clinical question:** {spec.get('clinical_question', 'not specified')}\n")
    lanes = spec.get("concept_lanes") or []
    if lanes:
        lines.append("## Concept Map Lanes\n")
        for lane in lanes:
            if isinstance(lane, dict):
                lines.append(f"- **{safe_text(lane.get('label', '')).title()}:** {lane.get('description', '')}")
            else:
                lines.append(f"- {safe_text(lane)}")
        lines.append("")
    lines.append("## Slide Notes\n")
    for i, slide in enumerate(slides, start=1):
        title = get_slide_value(slide, "title", f"slide {i}")
        learner_task = get_slide_value(slide, "learner_task", "")
        lines.append(f"- **{i}. {title}**")
        if learner_task:
            lines.append(f"  - Task: {learner_task}")
    questions = []
    for slide in slides:
        if normalize_key(slide.get("type", "")) == "retrieval_check":
            questions.extend(get_slide_value(slide, "questions", []) or [])
    if questions:
        lines.append("\n## Exit Check\n")
        for q in questions:
            lines.append(f"- {safe_text(q)}")
    out_path.write_text("\n".join(lines), encoding="utf-8")


def write_manifest(spec: Dict[str, Any], out_path: Path, files: List[str]) -> None:
    slides = spec.get("slides") or []
    metadata = spec.get("metadata") or {}
    slide_rows = []
    for i, slide in enumerate(slides, start=1):
        slide_rows.append({
            "slide_id": get_slide_value(slide, "slide_id", f"S{i:02d}"),
            "slide_number": i,
            "slide_title": get_slide_value(slide, "title", f"slide {i}"),
            "slide_type": slide.get("type", "generic"),
            "concept_lane": get_slide_value(slide, "concept_lane", ""),
            "concepts": get_slide_value(slide, "concepts", []) or [],
            "cjm_steps": get_slide_value(slide, "cjm_steps", []) or [],
            "learner_task": get_slide_value(slide, "learner_task", ""),
            "source_status": get_slide_value(slide, "source_status", metadata.get("source_status", "provisional")),
            "remediation_target": get_slide_value(slide, "remediation_target", ""),
        })
    coverage = cjm_coverage(slides)
    missing_cjm = [step for step in CJM_STEPS if coverage.get(step, 0) == 0]
    manifest = {
        "generated_at": dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z"),
        "package_id": metadata.get("package_id", ""),
        "package_status": spec.get("package_status", "draft-only"),
        "metadata": metadata,
        "runtime_config": spec.get("runtime_config", {}),
        "clinical_question": spec.get("clinical_question"),
        "slide_count": len(slides),
        "slides": slide_rows,
        "cjm_coverage": coverage,
        "files": files,
        "assumptions": spec.get("assumptions", []),
        "unresolved_variables": spec.get("unresolved_variables", []),
        "qa_summary": {
            "missing_cjm_functions": missing_cjm,
            "source_status": metadata.get("source_status", "provisional"),
            "notes": spec.get("qa_notes", []),
        },
        "revision_log": spec.get("revision_log", []),
    }
    out_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")


def build_default_slides(spec: Dict[str, Any]) -> List[Dict[str, Any]]:
    raise ValueError("default/demo slides are disabled for source-based builds; use --demo for a clearly labeled test package")


def make_demo_spec() -> Dict[str, Any]:
    return {
        "metadata": {
            "course": "maternal newborn nursing",
            "chapter": "chapter 2 active learning",
            "topic": "from biology to bedside",
            "subtitle": "reproductive anatomy, conception, fetal development, and early pregnancy safety decisions",
            "audience": "prelicensure nursing students",
            "duration_minutes": 75,
        },
        "clinical_question": "how does the body create, support, protect, and signal risk in early pregnancy?",
        "concept_lanes": [
            {"label": "create", "description": "anatomy and gametes make conception possible"},
            {"label": "support", "description": "hormones, placenta, membranes, cord, and circulation support growth"},
            {"label": "protect", "description": "timing and barriers shape preventable exposure risk"},
            {"label": "signal", "description": "bleeding, pain, fever, symptoms, and abnormal cues require action"},
        ],
        "quiz_analysis": [
            {
                "objective": "sort early pregnancy cues by urgency",
                "score_percent": 58,
                "miss_type": "risk priority",
                "cjm_step": "prioritize hypotheses",
                "remediation": "urgency sort plus patient teaching script",
            }
        ],
        "slides": [
            {
                "type": "title",
                "module_label": "chapter 2 active learning",
                "title": "From biology to bedside",
                "subtitle": "reproductive anatomy, conception, fetal development, and early pregnancy safety decisions",
                "cjm_steps": ["orientation"],
                "goals": ["build the map", "connect structure to cue", "practice safe teaching", "choose next action"],
            },
            {
                "type": "opening_case",
                "module_label": "opening case",
                "title": "Maya has a question",
                "patient_prompt": "I had a positive home pregnancy test. I am not sure what is normal. I had slight spotting yesterday.",
                "know": ["positive home pregnancy test", "spotting yesterday", "early pregnancy concern"],
                "need": ["bleeding amount", "pain or cramping", "dizziness/syncope", "fever", "gestational age or LMP"],
                "task": "Separate reassurance, assessment, teaching, and urgent cues. Do not jump to certainty.",
                "cjm_steps": ["recognize cues"],
                "concepts": ["signal"],
                "learner_task": "name known cues and missing safety data",
            },
            {
                "type": "clinical_question",
                "module_label": "organizing question",
                "title": "The clinical question that organizes the chapter",
                "cjm_steps": ["recognize cues", "analyze cues"],
                "concepts": ["create", "support", "protect", "signal"],
                "learner_task": "explain which lane a patient cue belongs to",
            },
            {
                "type": "chapter_map",
                "module_label": "chapter map",
                "title": "Chapter map: follow the patient question",
                "milestones": [
                    {"label": "anatomy", "body": "what structures matter?"},
                    {"label": "hormones", "body": "how is timing controlled?"},
                    {"label": "conception", "body": "where do cells meet and implant?"},
                    {"label": "fetal support", "body": "how does exchange occur?"},
                    {"label": "safety", "body": "what should the nurse teach or escalate?"},
                ],
                "cjm_steps": ["analyze cues"],
                "concepts": ["create", "support", "signal"],
                "learner_task": "use the map to locate a weak concept",
            },
            {
                "type": "warmup_sequence",
                "module_label": "warm-up",
                "title": "Warm-up: what must happen first?",
                "prompt": "Place these events in the most logical sequence.",
                "items": ["implantation in the endometrium", "fertilization in the fallopian tube", "early cell division during travel toward the uterus", "placental development begins"],
                "cjm_steps": ["analyze cues"],
                "concepts": ["create"],
                "learner_task": "sequence conception to implantation events",
            },
            {
                "type": "concept_cards",
                "module_label": "anatomy",
                "title": "Anatomy: structure drives the clinical question",
                "cards": [
                    {"title": "ovaries", "body": "produce ova and reproductive hormones"},
                    {"title": "fallopian tubes", "body": "site of fertilization and transport to uterus"},
                    {"title": "uterus", "body": "supports implantation, fetal growth, and contractions"},
                    {"title": "cervix and vagina", "body": "protect passage, secretions, and birth canal functions"},
                    {"title": "breasts", "body": "prepare for lactation through hormonal influence"},
                ],
                "callout": "When a structure matters, ask: what does it help happen or prevent?",
                "cjm_steps": ["recognize cues", "analyze cues"],
                "concepts": ["create", "support"],
                "learner_task": "connect each structure to a clinical function",
            },
            {
                "type": "match_activity",
                "module_label": "activity",
                "title": "Activity: match the structure to the cue",
                "sources": ["conception can happen before the uterus is involved", "heavy bleeding and cramping soon after early pregnancy", "cyclic dates help estimate pregnancy timing", "privacy concerns arise during reproductive assessment"],
                "targets": ["fallopian tubes", "uterus", "ovaries", "cervix / vagina / external structures"],
                "cjm_steps": ["recognize cues", "analyze cues"],
                "concepts": ["create", "signal"],
                "learner_task": "match a cue to the structure most directly involved",
            },
            {
                "type": "debrief",
                "module_label": "debrief",
                "title": "Debrief: structure-function reasoning",
                "cards": [
                    {"title": "fallopian tubes", "body": "fertilization usually occurs before the conceptus reaches the uterus."},
                    {"title": "uterus", "body": "heavy bleeding and cramping can signal a pregnancy complication and require assessment."},
                    {"title": "ovaries", "body": "ovarian and menstrual cycle timing helps estimate ovulation and gestational age."},
                    {"title": "cervical/vaginal/external structures", "body": "assessment requires consent, privacy, trauma-informed communication, and respectful terminology."},
                ],
                "cjm_steps": ["analyze cues"],
                "concepts": ["create", "signal"],
                "learner_task": "explain one match in structure-function language",
            },
            {
                "type": "control_room",
                "module_label": "hormones",
                "title": "Hormones: think control room, not flashcards",
                "center": "brain + gonads coordinate timing",
                "nodes": [
                    {"title": "hypothalamus", "body": "starts regulatory signaling"},
                    {"title": "ovaries / testes", "body": "produce gametes and reproductive hormones"},
                    {"title": "pituitary", "body": "FSH and LH support follicle development and ovulation"},
                    {"title": "estrogen / progesterone", "body": "prepare and maintain reproductive environment"},
                ],
                "cjm_steps": ["analyze cues"],
                "concepts": ["support"],
                "learner_task": "explain hormone control without memorizing isolated labels",
            },
            {
                "type": "timeline",
                "module_label": "cycle",
                "title": "Cycle timeline: fertility timing has a physiology story",
                "milestones": [
                    {"label": "follicular", "body": "follicle development and estrogen change"},
                    {"label": "ovulation", "body": "egg release after LH surge"},
                    {"label": "luteal", "body": "progesterone supports endometrium"},
                    {"label": "no implantation", "body": "hormone change and menstruation"},
                ],
                "cjm_steps": ["analyze cues"],
                "concepts": ["create", "support"],
                "learner_task": "connect timing to fertility and pregnancy testing questions",
            },
            {
                "type": "checkpoint_mcq",
                "module_label": "checkpoint",
                "title": "Checkpoint: choose the best explanation",
                "question": "A patient asks, 'Why does my cycle date matter if I already have a positive pregnancy test?' Best nursing response?",
                "options": [
                    "Cycle dates are only used for infertility, not pregnancy.",
                    "Cycle dates help estimate timing, but symptoms and prenatal follow-up still matter.",
                    "The exact conception day can always be calculated from the positive test.",
                    "You do not need to discuss cycle dates once pregnancy is confirmed.",
                ],
                "answer": "B",
                "rationale": "cycle timing supports estimation, but does not replace assessment and follow-up.",
                "cjm_steps": ["analyze cues", "generate solutions"],
                "concepts": ["support", "signal"],
                "learner_task": "select a patient-centered explanation",
            },
            {
                "type": "process_map",
                "module_label": "conception",
                "title": "Conception: from gametes to zygote",
                "milestones": [
                    {"label": "gametes", "body": "sperm and ovum carry genetic material"},
                    {"label": "fertilization", "body": "usually in fallopian tube"},
                    {"label": "zygote", "body": "chromosomal combination established"},
                    {"label": "cell division", "body": "conceptus travels toward uterus"},
                    {"label": "implantation", "body": "developing blastocyst embeds in endometrium"},
                ],
                "cjm_steps": ["analyze cues"],
                "concepts": ["create"],
                "learner_task": "map conception events to anatomic locations",
            },
            {
                "type": "teaching_point",
                "module_label": "implantation",
                "title": "Implantation: the point where early teaching becomes concrete",
                "cards": [
                    {"title": "blastocyst reaches uterus", "body": "developing structure reaches endometrium"},
                    {"title": "implants in endometrium", "body": "attachment begins pregnancy-supporting relationship"},
                    {"title": "placental development begins", "body": "early support structures start forming"},
                ],
                "warning": "The nurse explains expected possibilities and screens for urgent symptoms, not guarantees safety.",
                "cjm_steps": ["generate solutions", "take action"],
                "concepts": ["create", "signal"],
                "learner_task": "turn implantation science into safe patient teaching",
            },
            {
                "type": "exchange_model",
                "module_label": "placenta",
                "title": "Placenta: exchange organ, not a shield",
                "left": {"title": "maternal side", "body": "maternal blood supply"},
                "center": {"title": "placenta", "body": "oxygen + nutrients + waste exchange; hormones that support pregnancy; interface, not complete barrier"},
                "right": {"title": "fetal side", "body": "fetal circulation"},
                "warning": "The placenta supports exchange. It does not block every harmful exposure.",
                "cjm_steps": ["analyze cues", "prioritize hypotheses"],
                "concepts": ["support", "protect"],
                "learner_task": "explain why placenta function affects teaching about exposure risk",
            },
            {
                "type": "triad",
                "module_label": "support structures",
                "title": "Fetal support triad: fluid, membranes, cord",
                "cards": [
                    {"title": "amniotic fluid", "body": "cushions the fetus and supports movement; volume can affect assessment"},
                    {"title": "membranes", "body": "help protect the intrauterine environment; rupture changes infection risk"},
                    {"title": "umbilical cord", "body": "connects fetus and placenta; compression or compromise affects oxygenation"},
                ],
                "cjm_steps": ["recognize cues", "analyze cues"],
                "concepts": ["support", "protect"],
                "learner_task": "name the function and risk signal for each support structure",
            },
            {
                "type": "safety_chain",
                "module_label": "circulation",
                "title": "Fetal circulation: oxygenation depends on placenta and cord",
                "milestones": [
                    {"label": "placenta", "body": "gas exchange occurs here"},
                    {"label": "cord", "body": "connects fetus to placenta"},
                    {"label": "shunts", "body": "support fetal circulation pathways"},
                    {"label": "delivery", "body": "transition after birth changes flow"},
                    {"label": "monitor", "body": "later fetal assessment builds on this"},
                ],
                "cjm_steps": ["analyze cues", "prioritize hypotheses"],
                "concepts": ["support", "signal"],
                "learner_task": "connect support structures to fetal oxygenation risk",
            },
            {
                "type": "compare",
                "module_label": "development",
                "title": "Embryonic vs fetal period: timing changes risk",
                "columns": [
                    {"title": "embryonic period", "items": ["early organ development", "high vulnerability to exposures that affect organ formation"]},
                    {"title": "fetal period", "items": ["growth and maturation", "functional development and late-pregnancy growth concerns"]},
                ],
                "cjm_steps": ["prioritize hypotheses"],
                "concepts": ["protect"],
                "learner_task": "explain why timing changes exposure risk",
            },
            {
                "type": "risk_engine",
                "module_label": "teratogens",
                "title": "Teratogen risk engine: do not reduce risk to one variable",
                "center": "exposure risk question",
                "factors": [
                    {"title": "susceptibility", "body": "maternal and fetal factors"},
                    {"title": "duration", "body": "how long or how often"},
                    {"title": "timing", "body": "when during development"},
                    {"title": "dose", "body": "how much exposure"},
                    {"title": "context", "body": "medications, infection, environment, resources"},
                ],
                "cjm_steps": ["prioritize hypotheses", "generate solutions"],
                "concepts": ["protect", "signal"],
                "learner_task": "identify risk modifiers before giving teaching",
            },
            {
                "type": "mini_case",
                "module_label": "mini case",
                "title": "Mini case: respond to Maya without guessing",
                "scenario": "Maya says, 'I took ibuprofen twice and had alcohol before I knew. I feel okay today. Plus, I had light spotting yesterday.'",
                "steps": [
                    {"title": "1. assess first", "body": "bleeding amount, pain, dizziness, fever, LMP, medication dose, exposure details"},
                    {"title": "2. teach clearly", "body": "stop avoidable exposures now; review meds with prenatal provider; keep follow-up"},
                    {"title": "3. avoid false certainty", "body": "do not promise no risk; do not create panic; focus on urgent cues and next steps"},
                ],
                "cjm_steps": ["recognize cues", "generate solutions", "take action"],
                "concepts": ["protect", "signal"],
                "learner_task": "write one assessment sentence, one teaching sentence, and one safety limit",
            },
            {
                "type": "urgency_sort",
                "module_label": "clinical judgment",
                "title": "Clinical judgment: sort cues by urgency",
                "categories": [
                    {"title": "escalate now", "items": ["heavy bleeding", "severe abdominal pain", "syncope", "fever with worsening symptoms"]},
                    {"title": "assess and teach", "items": ["light spotting without red flags", "medication question", "mild expected symptom"]},
                ],
                "cjm_steps": ["prioritize hypotheses"],
                "concepts": ["signal"],
                "learner_task": "sort cues by immediate risk and explain why",
            },
            {
                "type": "script_template",
                "module_label": "practice",
                "title": "Practice: patient teaching script",
                "steps": [
                    {"title": "acknowledge", "body": "I hear that you are worried."},
                    {"title": "assess safety", "body": "ask targeted symptom and exposure questions."},
                    {"title": "explain simply", "body": "tie teaching to support, protection, and risk."},
                    {"title": "reduce risk", "body": "avoid preventable exposures and know urgent symptoms."},
                    {"title": "follow up", "body": "connect to prenatal care or escalation based on symptoms."},
                ],
                "cjm_steps": ["generate solutions", "take action"],
                "concepts": ["protect", "signal"],
                "learner_task": "build a patient-centered response using the template",
            },
            {
                "type": "basics_grid",
                "module_label": "genetics",
                "title": "Genetic basics: know your nursing lane",
                "cards": [
                    {"title": "what students should know", "body": "chromosomes carry genetic information; inheritance patterns can affect screening and risk discussion"},
                    {"title": "what nurses assess", "body": "family history, pregnancy history, congenital concerns, screening questions, patient understanding"},
                    {"title": "what nurses do not overclaim", "body": "do not provide genetic counseling beyond scope; refer to genetics resources when appropriate"},
                    {"title": "what nurses teach", "body": "plain language, consent, limits of screening, and follow-up"},
                ],
                "cjm_steps": ["analyze cues", "generate solutions"],
                "concepts": ["protect"],
                "learner_task": "separate basic knowledge from nursing scope and referral needs",
            },
            {
                "type": "capstone_mcq",
                "module_label": "capstone",
                "title": "Capstone scenario: choose the next safe action",
                "scenario": "Maya calls back: 'The spotting is heavier. I have sharp pain on one side and I almost fainted in the bathroom.'",
                "question": "What is the priority nursing response?",
                "options": [
                    "Explain that spotting can be normal and schedule routine teaching.",
                    "Ask her to track symptoms for 24 hours and call back if they continue.",
                    "Recognize urgent cues, assess immediate safety, and direct urgent evaluation/escalation per protocol.",
                    "Begin a detailed explanation of fetal development to reduce anxiety.",
                ],
                "answer": "C",
                "rationale": "heavier bleeding, unilateral sharp pain, and near-syncope are urgent cues requiring escalation, not reassurance alone.",
                "cjm_steps": ["prioritize hypotheses", "take action"],
                "concepts": ["signal"],
                "learner_task": "choose the safest next action and justify it",
            },
            {
                "type": "debrief_three",
                "module_label": "rationale",
                "title": "Capstone debrief: why C is safest",
                "cards": [
                    {"title": "recognize cues", "body": "heavier bleeding, sharp unilateral pain, near-syncope are not routine teaching cues."},
                    {"title": "analyze risk", "body": "early pregnancy complication is possible; maternal safety comes first."},
                    {"title": "take action", "body": "escalate according to protocol and direct urgent evaluation; do not reassure without assessment."},
                ],
                "cjm_steps": ["recognize cues", "analyze cues", "take action", "evaluate outcomes"],
                "concepts": ["signal"],
                "learner_task": "explain why the distractors are unsafe or incomplete",
            },
            {
                "type": "retrieval_check",
                "module_label": "rapid retrieval",
                "title": "Rapid retrieval: five-item exit check",
                "questions": [
                    "Where does fertilization usually occur?",
                    "What hormone surge is associated with ovulation?",
                    "What organ supports exchange but is not a complete barrier?",
                    "Name one urgent early pregnancy cue.",
                    "What are the four nursing action frames from this lesson?",
                ],
                "cjm_steps": ["evaluate outcomes"],
                "concepts": ["create", "support", "protect", "signal"],
                "learner_task": "answer by spoken or written exit check",
            },
            {
                "type": "takeaway",
                "module_label": "takeaway",
                "title": "Chapter 2 takeaway: biology becomes nursing action",
                "actions": [
                    {"title": "teach", "body": "use plain language for timing, anatomy, implantation, and fetal support"},
                    {"title": "assess", "body": "ask about history, symptoms, medications, exposures, and risk factors"},
                    {"title": "escalate", "body": "prioritize bleeding, severe pain, syncope, fever, and abnormal cues"},
                    {"title": "reassure", "body": "only after understanding, context, and red flags are ruled out"},
                ],
                "statement": "A safe nurse can explain the process, reduce preventable risk, screen for urgent cues, and connect the patient to timely follow-up.",
                "cjm_steps": ["take action", "evaluate outcomes"],
                "concepts": ["create", "support", "protect", "signal"],
                "learner_task": "state one bedside action for each frame",
            },
        ],
    }


def load_spec(args) -> Dict[str, Any]:
    if args.demo:
        return make_demo_spec()
    if not args.spec:
        raise SystemExit("Provide --spec lesson_spec.json or use --demo.")
    path = Path(args.spec)
    if not path.exists():
        raise SystemExit(f"Spec not found: {path}")
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise SystemExit(f"Invalid JSON in spec: {exc}") from exc


def generate_package(spec: Dict[str, Any], outdir: Path) -> Dict[str, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    deck_path = outdir / "lesson_deck.pptx"
    facilitator_path = outdir / "facilitator_guide.md"
    learner_path = outdir / "learner_handout.md"
    map_path = outdir / "assessment_map.csv"
    manifest_path = outdir / "lesson_manifest.json"

    make_presentation(spec, deck_path)
    write_assessment_map(spec, map_path)
    write_facilitator_guide(spec, facilitator_path)
    write_learner_handout(spec, learner_path)
    write_manifest(
        spec,
        manifest_path,
        [deck_path.name, facilitator_path.name, learner_path.name, map_path.name, manifest_path.name],
    )
    return {
        "deck": str(deck_path),
        "facilitator_guide": str(facilitator_path),
        "learner_handout": str(learner_path),
        "assessment_map": str(map_path),
        "manifest": str(manifest_path),
    }


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="Generate a Harrity lesson package from a JSON spec.")
    parser.add_argument("--spec", help="Path to lesson_spec.json")
    parser.add_argument("--demo", action="store_true", help="Generate demo early-pregnancy lesson package")
    parser.add_argument("--outdir", default="harrity_lesson_package", help="Output directory")
    args = parser.parse_args(argv)

    spec = load_spec(args)
    outdir = Path(args.outdir)
    files = generate_package(spec, outdir)
    print(json.dumps(files, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
